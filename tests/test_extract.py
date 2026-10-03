"""从排好版的模板读取格式：用本工具按已知要求生成一份文档当模板，读回来应得到同样的要求。"""
from docx import Document

from gongwen_formatter import format_file
from gongwen_formatter.classify import Role
from gongwen_formatter.extract import extract
from gongwen_formatter.spec import parse

CUSTOM = """\
标题（黑体，小二，居中，行距固定值30磅）
正文（宋体，小四，首行缩进2字符，1.5倍行距，两端对齐）
一级标题（黑体，四号，顶格，1.5倍行距）
二级标题（楷体，小四，加粗，首行缩进2字符，1.5倍行距）
三级标题（宋体，小四，加粗，首行缩进2字符，1.5倍行距）
四级标题（楷体，小四，加粗，首行缩进2字符，1.5倍行距）
落款（宋体，小四，右对齐，右缩进2字符，1.5倍行距）
日期（楷体，小四，右对齐，右缩进2字符，1.5倍行距）
英文、数字：Arial
页边距：上25毫米，下25毫米，左30毫米，右20毫米
页眉15毫米，页脚28毫米
"""
LINES = ["关于开展志愿服务月活动的通知", "一、活动目的", "（一）弘扬志愿精神",
         "通过系列活动，提升学生社会责任感，计划覆盖2000人次。",
         "1. 时间安排：10月1日至10月31日。", "（1）报名：10月1日前。", "二、其他", "请各班按时报名。",
         "金融与统计学院", "2026年10月1日"]


def _make_template(tmp_path):
    src, tpl = tmp_path / "src.docx", tmp_path / "tpl.docx"
    d = Document()
    for line in LINES:
        d.add_paragraph(line)
    d.save(src)
    format_file(str(src), str(tpl), CUSTOM)
    return tpl


def test_roundtrip_requirements(tmp_path):
    tpl = _make_template(tmp_path)
    got = extract(str(tpl))
    want = parse(CUSTOM)
    back = parse(got.text)
    assert back.problems == []
    for role in (Role.TITLE, Role.BODY, Role.H1, Role.H2, Role.H3, Role.H4, Role.SIGNATURE, Role.DATE):
        a, b = want.styles[role], back.styles[role]
        assert (a.font, a.size, a.bold, a.align, a.first_indent, a.right_indent, a.line, a.line_multiple) == \
               (b.font, b.size, b.bold, b.align, b.first_indent, b.right_indent, b.line, b.line_multiple), role
    assert back.latin_font == "Arial"
    assert back.margins_mm == (25, 25, 30, 20)
    assert (back.header_mm, back.footer_mm) == (15, 28)


def test_missing_roles_follow_template_body(tmp_path):
    tpl = _make_template(tmp_path)
    got = extract(str(tpl))
    assert any("附件" in n and "推算" in n for n in got.notes)
    assert parse(got.text).styles[Role.ATTACH_LABEL].font == "宋体"


def test_unbold_h2_survives_template_roundtrip(tmp_path):
    src, tpl = tmp_path / "src.docx", tmp_path / "tpl.docx"
    doc = Document()
    for text in LINES:
        doc.add_paragraph(text)
    doc.save(src)
    format_file(src, tpl, "二级标题（楷体，小四，不加粗）")
    assert parse(extract(tpl).text).styles[Role.H2].bold is False
