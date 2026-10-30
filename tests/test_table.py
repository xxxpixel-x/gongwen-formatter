"""表格：默认保持原稿的样子（包括继承来的字体、样式里的边框和底纹）；写了“表头”“表格”时只改写到的项；
表格标题能识别；模板里的表格格式以注释写出。"""
from docx import Document
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor

from gongwen_formatter import format_file
from gongwen_formatter.classify import Role
from gongwen_formatter.extract import extract
from gongwen_formatter.render import read_document
from gongwen_formatter.spec import KEEP, describe_role, parse
from gongwen_formatter.table import cells, header_rows

ROWS = [["序号", "经费项目", "规格 / 说明", "小计（元）"],
        ["1", "教具制作费", "1、绘画纸、画笔", "1200"],
        ["2", "志愿者补贴", "(2)交通补贴", ""]]


def _source(path, style="Table Grid", width_mm=None):
    d = Document()
    d.add_paragraph("一、经费预算")
    t = d.add_table(rows=len(ROWS), cols=len(ROWS[0]))
    if style:
        t.style = d.styles[style]
    for r, values in enumerate(ROWS):
        for c, value in enumerate(values):
            run = t.cell(r, c).paragraphs[0].add_run(value)
            run.font.size, run.font.color.rgb = Pt(9), RGBColor(0x66, 0x66, 0x66)
            run.bold = c == 0
    if width_mm:
        for col in t.columns:
            for cell in col.cells:
                cell.width = Mm(width_mm)
        for g in t._tbl.tblGrid.findall(qn("w:gridCol")):
            g.set(qn("w:w"), str(round(width_mm * 1440 / 25.4)))
    d.add_paragraph("请各班按时报名。")
    d.save(path)


def _run_fmt(cell):
    rpr = [r for r in cell.paragraphs[0].runs if r.text][0]._r.rPr
    b = rpr.find(qn("w:b"))
    color = rpr.find(qn("w:color"))
    return {
        "font": rpr.find(qn("w:rFonts")).get(qn("w:eastAsia")),
        "size": int(rpr.find(qn("w:sz")).get(qn("w:val"))) / 2,
        "bold": b is not None and b.get(qn("w:val"), "1") not in ("0", "false"),
        "color": color.get(qn("w:val")) if color is not None else None,
    }


def _out(tmp_path, req=None, **kw):
    src, out = tmp_path / "in.docx", tmp_path / "out.docx"
    _source(src, **kw)
    format_file(str(src), str(out), req) if req else format_file(str(src), str(out))
    return Document(out).tables[0]


# ---------- 默认保持原样 ----------

def test_table_kept_as_is(tmp_path):
    table = _out(tmp_path)
    assert [[c.text for c in row.cells] for row in table.rows] == ROWS   # 编号写法也不改
    assert _run_fmt(table.cell(1, 0))["bold"] and not _run_fmt(table.cell(1, 1))["bold"]
    assert _run_fmt(table.cell(1, 1))["size"] == 9
    assert _run_fmt(table.cell(1, 1))["color"] == "666666"
    tbl = table._tbl
    assert tbl.find(".//" + qn("w:vAlign")) is None


def test_inherited_formats_written_on_cells(tmp_path):
    """原稿单元格没有直接设格式、靠“正文”样式和表格样式继承时，输出里也要一样，
    不能变成新文档的仿宋三号、28 磅行距。"""
    src, out = tmp_path / "in.docx", tmp_path / "out.docx"
    d = Document()
    normal = d.styles["Normal"]
    normal.font.size = Pt(12)
    normal.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "宋体")
    d.add_paragraph("一、经费预算")
    t = d.add_table(rows=2, cols=2)
    t.style = d.styles["Table Grid"]   # 样式里段后 0、单倍行距
    t.cell(0, 0).text, t.cell(1, 0).text = "项目", "教具"
    d.save(src)
    format_file(str(src), str(out))
    p = Document(out).tables[0].cell(1, 0).paragraphs[0]
    rpr = p.runs[0]._r.rPr
    assert rpr.find(qn("w:rFonts")).get(qn("w:eastAsia")) == "宋体"
    assert rpr.find(qn("w:sz")).get(qn("w:val")) == "24"
    spacing = p._p.pPr.find(qn("w:spacing"))
    assert (spacing.get(qn("w:after")), spacing.get(qn("w:line")), spacing.get(qn("w:lineRule"))) == \
        ("0", "240", "auto")
    assert p._p.pPr.find(qn("w:pStyle")) is None


def test_border_colors_kept(tmp_path):
    src, out = tmp_path / "in.docx", tmp_path / "out.docx"
    _source(src, style=None)
    d = Document(src)
    tcpr = d.tables[0].cell(0, 0)._tc.get_or_add_tcPr()
    tcb = tcpr.makeelement(qn("w:tcBorders"), {})
    for side, val in (("top", "single"), ("bottom", "none")):
        tcb.append(tcb.makeelement(qn("w:" + side), {qn("w:val"): val, qn("w:color"): "CCCCCC", qn("w:sz"): "4"}))
    tcpr.append(tcb)
    d.save(src)
    format_file(str(src), str(out))
    tcb = Document(out).tables[0]._tbl.find(".//" + qn("w:tcBorders"))
    assert tcb.find(qn("w:top")).get(qn("w:color")) == "CCCCCC"
    assert tcb.find(qn("w:bottom")).get(qn("w:val")) == "none"


def test_style_borders_kept(tmp_path):
    borders = _out(tmp_path)._tbl.tblPr.find(qn("w:tblBorders"))
    assert borders is not None, "“网格型”样式的边框应写到表格上，否则输出文档里边框会消失"
    assert {e.get(qn("w:val")) for e in borders} == {"single"}


def test_style_shading_and_text_color_kept(tmp_path):
    src, out = tmp_path / "in.docx", tmp_path / "out.docx"
    d = Document()
    d.add_paragraph("一、经费预算")
    t = d.add_table(rows=4, cols=2)
    t.style = d.styles["Medium Shading 1 Accent 1"]   # 深蓝表头白字 + 隔行浅蓝
    for r in range(4):
        t.cell(r, 0).text = f"第{r}行"
    t.cell(2, 1).paragraphs[0].add_run("自设")
    tcpr = t.cell(2, 1)._tc.get_or_add_tcPr()
    tcpr.append(tcpr.makeelement(qn("w:shd"), {qn("w:val"): "clear", qn("w:color"): "auto",
                                               qn("w:fill"): "FFFF00"}))
    d.save(src)
    format_file(str(src), str(out))
    table = Document(out).tables[0]
    fill = lambda c: (c._tc.tcPr.find(qn("w:shd")).get(qn("w:fill"))
                      if c._tc.tcPr.find(qn("w:shd")) is not None else None)
    assert [fill(table.cell(r, 0)) for r in range(4)] == ["4F81BD", "D3DFEE", None, "D3DFEE"]
    assert fill(table.cell(2, 1)) == "FFFF00", "单元格自己设的底纹优先"
    head = _run_fmt(table.cell(0, 0))
    assert head["color"] == "FFFFFF" and head["bold"], "表头的白字、加粗来自表格样式，也要保留"
    tbl = next(i.raw for i in read_document(str(src)).items if i.raw is not None)
    first = {(c.row, c.col): c for c in cells(tbl)}[(0, 0)]
    assert (first.fill, first.color, first.bold) == ("4F81BD", "FFFFFF", True)


def test_wide_table_shrinks_to_text_width(tmp_path):
    tbl = _out(tmp_path, width_mm=60)._tbl   # 4 列 × 60 毫米，比版心（156 毫米）宽
    grid = [int(g.get(qn("w:w"))) for g in tbl.tblGrid.findall(qn("w:gridCol"))]
    assert sum(grid) <= round(156 * 1440 / 25.4)
    assert len(set(grid)) == 1


def test_original_file_untouched(tmp_path):
    src, out = tmp_path / "in.docx", tmp_path / "out.docx"
    _source(src)
    before = src.read_bytes()
    format_file(str(src), str(out))
    assert src.read_bytes() == before


# ---------- 写了“表头”“表格”时只改写到的项 ----------

def test_only_written_fields_change(tmp_path):
    table = _out(tmp_path, "表格：楷体\n表头（四号，不加粗）")
    body = _run_fmt(table.cell(1, 1))
    assert body["font"] == "楷体" and body["size"] == 9 and body["color"] == "666666"
    assert _run_fmt(table.cell(1, 0))["bold"], "没写加粗，原稿的加粗保留"
    head = _run_fmt(table.cell(0, 0))
    assert head["size"] == 14 and not head["bold"] and head["font"] != "楷体"


def test_table_line_spacing_and_alignment(tmp_path):
    table = _out(tmp_path, "表格（左对齐，行距固定值20磅）")
    ppr = table.cell(1, 1).paragraphs[0]._p.pPr
    assert ppr.find(qn("w:jc")).get(qn("w:val")) == "left"
    spacing = ppr.find(qn("w:spacing"))
    assert (spacing.get(qn("w:line")), spacing.get(qn("w:lineRule"))) == ("400", "exact")
    assert _run_fmt(table.cell(1, 1))["size"] == 9


def test_parse_table_fields():
    spec = parse("表头：宋体，小三，加粗\n表格内容（楷体）")
    assert spec.table_fields == {Role.TABLE_HEAD: {"font", "size", "bold"}, Role.TABLE: {"font"}}
    assert not spec.problems
    shown = describe_role(spec, Role.TABLE)
    assert shown["字体"] == "楷体" and shown["字号"] == KEEP
    default = parse("")
    assert default.table_fields == {}
    assert set(describe_role(default, Role.TABLE_HEAD).values()) == {KEEP}


# ---------- 结构 ----------

def test_header_rows_and_merged_cells():
    d = Document()
    t = d.add_table(rows=3, cols=3)
    t.cell(0, 0).merge(t.cell(0, 1))
    t.cell(1, 2).merge(t.cell(2, 2))
    t.cell(0, 0).text, t.cell(1, 2).text = "合并标题", "纵向"
    tbl = t._tbl
    assert header_rows(tbl) == {0}
    got = {(c.row, c.col): (c.colspan, c.rowspan, c.head) for c in cells(tbl)}
    assert got[(0, 0)] == (2, 1, True)
    assert got[(1, 2)] == (1, 2, False)
    assert (2, 2) not in got

    # 设了“标题行重复”的行优先；单行表格没有表头
    for tr in tbl.tr_lst[:2]:
        tr.get_or_add_trPr().append(tr.makeelement(qn("w:tblHeader"), {}))
    assert header_rows(tbl) == {0, 1}
    single = d.add_table(rows=1, cols=2)._tbl
    assert header_rows(single) == set()


def test_read_document_lists_table(tmp_path):
    src = tmp_path / "in.docx"
    _source(src)
    result = read_document(str(src))
    table = [i for i in result.items if i.raw is not None]
    assert len(table) == 1 and table[0].text == "[表格 3行×4列]"
    assert not any("表格里的编号" in w for w in result.warnings)


# ---------- 模板 ----------

def test_extract_writes_table_formats_as_comments(tmp_path):
    src, tpl = tmp_path / "in.docx", tmp_path / "tpl.docx"
    _source(src)
    req = "表头（宋体，四号，加粗，居中，顶格，单倍行距）\n表格（楷体，五号，不加粗，左对齐，顶格，单倍行距）"
    format_file(str(src), str(tpl), req)
    got = extract(str(tpl))
    assert parse(got.text).table_fields == {}, "读模板不能让表格自动被改"
    enabled = "\n".join(line[2:] if line.startswith("# 表") else line for line in got.text.splitlines())
    spec, want = parse(enabled), parse(req)
    for role in (Role.TABLE_HEAD, Role.TABLE):
        assert spec.styles[role] == want.styles[role]


def test_extract_without_table_has_no_table_lines(tmp_path):
    src, tpl = tmp_path / "in.docx", tmp_path / "tpl.docx"
    d = Document()
    d.add_paragraph("一、活动目的")
    d.add_paragraph("请各班按时报名。")
    d.save(src)
    format_file(str(src), str(tpl))
    assert "表头" not in extract(str(tpl)).text


# ---------- 表格标题 ----------

def _with_caption(path, before=None, after=None, first="一、经费预算"):
    d = Document()
    d.add_paragraph(first)
    if before is not None:
        d.add_paragraph(before)
        d.add_paragraph("")   # 中间隔一个空段也算紧挨着
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text, t.cell(1, 0).text = "项目", "1"
    if after is not None:
        d.add_paragraph(after)
    d.add_paragraph("请各班按时报名。")
    d.save(path)


def _roles(path):
    return {it.text: it for it in read_document(str(path)).items}


def test_numbered_caption_above_or_below_table(tmp_path):
    for before, after in (("表1 经费明细", None), (None, "续表 2-1：人员安排"), ("表一　活动安排", None)):
        src = tmp_path / "in.docx"
        _with_caption(src, before, after)
        text = before or after
        it = _roles(src)[text.replace("　", "")]
        assert it.role == Role.TABLE_CAPTION and it.confidence == "high", text


def test_named_table_above_is_low_confidence_caption(tmp_path):
    src = tmp_path / "in.docx"
    _with_caption(src, "志愿服务经费明细表")
    it = _roles(src)["志愿服务经费明细表"]
    assert it.role == Role.TABLE_CAPTION and it.confidence == "low"


def test_not_caption(tmp_path):
    src = tmp_path / "in.docx"
    _with_caption(src, "表彰大会安排")            # “表”后不是编号
    assert _roles(src)["表彰大会安排"].role != Role.TABLE_CAPTION
    _with_caption(src, None, None)
    d = Document(src)
    d.add_paragraph("表1 不挨着表格")
    d.save(src)
    assert _roles(src)["表1 不挨着表格"].role != Role.TABLE_CAPTION


def test_caption_formatting_and_numbering_check(tmp_path):
    src, out = tmp_path / "in.docx", tmp_path / "out.docx"
    _with_caption(src, "表1 经费明细")
    assert not any("表1" in w for w in read_document(str(src)).warnings)
    format_file(str(src), str(out), "表格标题（宋体，四号，加粗）")
    p = next(p for p in Document(out).paragraphs if p.text == "表1 经费明细")
    rpr = p.runs[0]._r.rPr
    assert rpr.find(qn("w:rFonts")).get(qn("w:eastAsia")) == "宋体"
    assert rpr.find(qn("w:b")) is not None
    assert p._p.pPr.find(qn("w:jc")).get(qn("w:val")) == "center"


def test_extract_reads_caption_format(tmp_path):
    src, tpl = tmp_path / "in.docx", tmp_path / "tpl.docx"
    _with_caption(src, "表1 经费明细")
    req = "表格标题（宋体，四号，加粗，居中，顶格，单倍行距）"
    format_file(str(src), str(tpl), req)
    got = extract(str(tpl))
    assert parse(got.text).styles[Role.TABLE_CAPTION] == parse(req).styles[Role.TABLE_CAPTION]
