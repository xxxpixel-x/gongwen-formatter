"""读写失败和路径别名不能损坏原稿或已有输出。"""
import importlib
import os
import struct
import zlib
from io import BytesIO

import pytest
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.opc.packuri import PackURI
from docx.opc.part import Part

from gongwen_formatter import format_file, read_document
from gongwen_formatter.cli import main

render_module = importlib.import_module("gongwen_formatter.render")


def make_source(path):
    doc = Document()
    doc.add_paragraph("关于开展工作的通知")
    doc.add_paragraph("请各单位按时提交材料。")
    doc.save(path)


@pytest.mark.parametrize("alias", ["same", "symlink", "hardlink"])
def test_api_refuses_source_alias(tmp_path, alias):
    source = tmp_path / "原稿.docx"
    make_source(source)
    original = source.read_bytes()
    target = source if alias == "same" else tmp_path / "输出.docx"
    if alias != "same":
        try:
            os.symlink(source, target) if alias == "symlink" else os.link(source, target)
        except OSError:
            pytest.skip("当前平台不允许创建此类链接")
    with pytest.raises(ValueError, match="原稿"):
        format_file(source, target)
    assert source.read_bytes() == original


def test_cli_refuses_source_path(tmp_path, capsys):
    source = tmp_path / "原稿.docx"
    make_source(source)
    original = source.read_bytes()
    assert main([str(source), "-o", str(source)]) != 0
    assert source.read_bytes() == original
    assert "原稿" in capsys.readouterr().err


@pytest.mark.parametrize("kind", ["missing", "invalid_zip"])
def test_cli_reports_bad_input_without_traceback(tmp_path, capsys, kind):
    source = tmp_path / "原稿.docx"
    if kind == "invalid_zip":
        source.write_text("这不是 Word 文件", encoding="utf-8")
    assert main([str(source)]) != 0
    assert "读取失败" in capsys.readouterr().err
    assert not (tmp_path / "原稿_已排版.docx").exists()


def test_failed_save_preserves_existing_output(tmp_path, monkeypatch):
    source, target = tmp_path / "原稿.docx", tmp_path / "输出.docx"
    make_source(source)
    make_source(target)
    original, old_output = source.read_bytes(), target.read_bytes()

    def broken_save(self, path):
        with open(path, "wb") as stream:
            stream.write(b"partial zip")
        raise OSError("模拟磁盘写入失败")

    monkeypatch.setattr("docx.document.Document.save", broken_save)
    with pytest.raises(OSError, match="模拟"):
        format_file(source, target)
    assert source.read_bytes() == original
    assert target.read_bytes() == old_output
    assert set(tmp_path.iterdir()) == {source, target}


def test_failed_replace_preserves_existing_output(tmp_path, monkeypatch):
    source, target = tmp_path / "原稿.docx", tmp_path / "输出.docx"
    make_source(source)
    make_source(target)
    previous = target.read_bytes()

    def locked(*args):
        raise PermissionError("文件正在使用")

    monkeypatch.setattr(render_module.os, "replace", locked)
    with pytest.raises(PermissionError):
        format_file(source, target)
    assert target.read_bytes() == previous
    assert set(tmp_path.iterdir()) == {source, target}


def _numbering(parent, num_id):
    numpr = OxmlElement("w:numPr")
    element = OxmlElement("w:numId")
    element.set(qn("w:val"), str(num_id))
    numpr.append(element)
    parent.append(numpr)


@pytest.mark.parametrize("disabled", [False, True])
def test_inherited_automatic_numbering_warning(tmp_path, disabled):
    path = tmp_path / "原稿.docx"
    doc = Document()
    doc.add_paragraph("通知")
    base = doc.styles.add_style("自动编号基样式", WD_STYLE_TYPE.PARAGRAPH)
    _numbering(base.element.get_or_add_pPr(), 1)
    child = doc.styles.add_style("继承编号", WD_STYLE_TYPE.PARAGRAPH)
    child.base_style = base
    paragraph = doc.add_paragraph("需处理的内容。", style=child)
    if disabled:
        _numbering(paragraph._p.get_or_add_pPr(), 0)
    doc.save(path)
    warnings = read_document(path).warnings
    assert any("自动编号" in w for w in warnings) is not disabled


def test_table_hyperlink_survives_formatting(tmp_path):
    source, target = tmp_path / "原稿.docx", tmp_path / "输出.docx"
    doc = Document()
    doc.add_paragraph("通知")
    paragraph = doc.add_table(rows=1, cols=1).cell(0, 0).paragraphs[0]
    rid = doc.part.relate_to("https://example.com/instructions", RT.HYPERLINK, is_external=True)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), rid)
    run, text = OxmlElement("w:r"), OxmlElement("w:t")
    text.text = "办理说明"
    run.append(text)
    link.append(run)
    paragraph._p.append(link)
    doc.save(source)
    format_file(source, target)
    result = Document(target)
    hyperlink = result.tables[0]._tbl.find(".//" + qn("w:hyperlink"))
    relationship = result.part.rels[hyperlink.get(qn("r:id"))]
    assert relationship.is_external
    assert relationship.target_ref == "https://example.com/instructions"


def _png(rgb):
    def chunk(kind, data):
        return struct.pack(">I", len(data)) + kind + data + struct.pack(">I", zlib.crc32(kind + data))
    return (b"\x89PNG\r\n\x1a\n" + chunk(b"IHDR", struct.pack(">IIBBBBB", 1, 1, 8, 2, 0, 0, 0))
            + chunk(b"IDAT", zlib.compress(b"\0" + bytes(rgb))) + chunk(b"IEND", b""))


def test_table_images_keep_correct_content_and_shared_relationships(tmp_path):
    source, target = tmp_path / "原稿.docx", tmp_path / "输出.docx"
    doc = Document()
    doc.add_paragraph("通知")
    images = [_png((255, 0, 0)), _png((0, 255, 0))]
    # 正文中的图片会跳过；表格重复使用其中一张图片，并包含另一张，不能串图。
    doc.add_paragraph().add_run().add_picture(BytesIO(images[0]))
    for image in [*images, images[0]]:
        cell = doc.add_table(rows=1, cols=1).cell(0, 0)
        cell.paragraphs[0].add_run().add_picture(BytesIO(image))
    doc.save(source)
    format_file(source, target)
    result = Document(target)
    for table, expected in zip(result.tables, [*images, images[0]]):
        blip = table._tbl.find(".//" + qn("a:blip"))
        rel = result.part.rels[blip.get(qn("r:embed"))]
        assert rel.reltype == RT.IMAGE
        assert rel.target_part.blob == expected


@pytest.mark.parametrize("kind", ["chart", "invalid_image"])
def test_unsupported_table_resource_preserves_existing_output(tmp_path, kind):
    source, target = tmp_path / "原稿.docx", tmp_path / "输出.docx"
    doc = Document()
    doc.add_paragraph("通知")
    cell = doc.add_table(rows=1, cols=1).cell(0, 0)
    if kind == "chart":
        part = Part(PackURI("/word/charts/chart1.xml"),
                    "application/vnd.openxmlformats-officedocument.drawingml.chart+xml", b"<chart/>")
        rid = doc.part.relate_to(part, RT.CHART)
        element = OxmlElement("c:chart")
        element.set(qn("r:id"), rid)
        cell.paragraphs[0]._p.append(element)
    else:
        cell.paragraphs[0].add_run().add_picture(BytesIO(_png((0, 0, 0))))
        image_part = next(rel.target_part for rel in doc.part.rels.values() if rel.reltype == RT.IMAGE)
        image_part._blob = b"invalid image"
    doc.save(source)
    make_source(target)
    original, previous = source.read_bytes(), target.read_bytes()
    with pytest.raises(ValueError, match="表格"):
        format_file(source, target)
    assert target.read_bytes() == previous
    assert source.read_bytes() == original
    assert set(tmp_path.iterdir()) == {source, target}


def test_cli_reads_utf8_bom_requirements(tmp_path, capsys):
    source, requirements = tmp_path / "原稿.docx", tmp_path / "格式.txt"
    make_source(source)
    requirements.write_text("正文（宋体，小四，不加粗）", encoding="utf-8-sig")
    assert main([str(source), "-r", str(requirements), "--check"]) == 0
    output = capsys.readouterr().out
    assert "没看懂" not in output and "不可识别" not in output
    assert "宋体，小四" in output
