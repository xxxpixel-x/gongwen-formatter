"""端到端：乱格式原稿 → 按默认要求排版 → 检查每类段落的格式。"""
from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt

from gongwen_formatter import format_file

LINES = [
    "关于开展志愿服务月活动的通知", "一、活动目的", "（一）弘扬志愿精神",
    "　　通过系列活动，提升学生社会责任感，计划覆盖2000人次。",
    "1. 时间安排：10月1日至10月31日。", "相关说明", "请各班按时报名。",
    "附件：", "1. 报名表", "金融与统计学院", "2026年10月1日",
]


def _messy(path):
    d = Document()
    for i, line in enumerate(LINES):
        r = d.add_paragraph(line).runs[0]
        r.font.size, r.font.name, r.bold = Pt(9 + i), "微软雅黑", i % 2 == 0
        d.add_paragraph("")
    d.save(path)


def _fmt(p, run=0):
    r = [x for x in p.runs if x.text][run]._r.rPr
    f = r.find(qn("w:rFonts"))
    sp = p._p.pPr.find(qn("w:spacing"))
    return {
        "font": f.get(qn("w:eastAsia")), "latin": f.get(qn("w:ascii")),
        "size": int(r.find(qn("w:sz")).get(qn("w:val"))) / 2,
        "bold": r.find(qn("w:b")) is not None,
        "jc": p._p.pPr.find(qn("w:jc")).get(qn("w:val")),
        "indent": int(p._p.pPr.find(qn("w:ind")).get(qn("w:firstLineChars"))),
        "line": int(sp.get(qn("w:line"))) / 20,
    }


def test_default_spec_end_to_end(tmp_path):
    src, out = tmp_path / "in.docx", tmp_path / "out.docx"
    _messy(src)
    items, warnings = format_file(str(src), str(out))
    assert warnings == []

    paras = {p.text: p for p in Document(out).paragraphs if p.text}
    title = _fmt(paras[LINES[0]])
    assert title == dict(font="方正小标宋_GBK", latin="Times New Roman", size=22,
                         bold=False, jc="center", indent=0, line=28)
    assert _fmt(paras["一、活动目的"])["font"] == "黑体"
    assert _fmt(paras["（一）弘扬志愿精神"])["bold"] is True
    body = _fmt(paras[LINES[3].strip("　")])
    assert (body["font"], body["size"], body["indent"], body["jc"]) == ("仿宋_GB2312", 16, 200, "both")
    h3 = paras["1. 时间安排：10月1日至10月31日。"]
    assert _fmt(h3, 0)["bold"] and not _fmt(h3, 1)["bold"]   # 只加粗到冒号
    assert _fmt(paras["相关说明"])["font"] == "黑体"          # 无编号的小标题
    assert _fmt(paras["2026年10月1日"])["jc"] == "right"

    # 多余空行已删除，只在附件前、落款前各留一行
    blanks = [p for p in Document(out).paragraphs if not p.text]
    assert len(blanks) == 2

    sec = Document(out).sections[0]
    assert round(sec.top_margin.mm) == 37 and round(sec.left_margin.mm) == 28
