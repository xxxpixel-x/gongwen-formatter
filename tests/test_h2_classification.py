"""带二级编号的正文不能仅凭编号被排成楷体加粗；覆盖截图反馈及相邻边界。"""
import pytest
from docx import Document
from docx.oxml.ns import qn

from gongwen_formatter import format_file
from gongwen_formatter.classify import Role, check_numbering, classify, lead_len


CONCLUSION = [
    "（一）本规划是部门未来三年发展的行动纲领，全体人员须统一思想、坚定信心、真抓实干。",
    "（二）在执行过程中，将根据实际情况适时调整优化，确保规划科学可行、务实有效，为部门持续健康发展奠定坚实基础。",
]


def test_conclusion_numbered_paragraphs_are_body():
    items = classify(["部门三年发展规划", "一、结语", *CONCLUSION])
    for item, text in zip(items[2:], CONCLUSION):
        assert item.text == text
        assert item.role == Role.BODY
        assert item.lead_len == 0
        assert item.confidence == "low"
        assert not item.confirmed
        assert "正文" in item.note
    assert check_numbering(items) == []


@pytest.mark.parametrize("text", [
    *CONCLUSION,
    CONCLUSION[0].removesuffix("。"),  # 漏写句号也不应整段变成标题
    CONCLUSION[0] + "请各科室认真落实。",  # 不能把第一整句当作标题
    "（一）请按时提交材料。",
    "（一）请按时提交材料！",
    "（一）请按时提交材料；逾期不予受理。",
    "（一）本规划已经发布，各科室应认真落实。后续将开展检查。",
    "（一）请认真学习《工作方案》。",
    "（一）有关要求是“按时完成。”",
    "（一）",
])
def test_h2_body_candidates_are_reviewable(text):
    item = classify(["标题", text])[1]
    assert item.role == Role.BODY
    assert item.lead_len == 0
    assert item.confidence == "low"
    assert not item.confirmed


@pytest.mark.parametrize("text, heading", [
    ("（一）工作目标。今年完成三项任务。", "（一）工作目标。"),
    ("（一）工作目标：今年完成三项任务。", "（一）工作目标："),
    ("(一)、工作目标:今年完成三项任务。", "（一）工作目标:"),
    ("（一）统筹推进重点项目建设和公共服务保障工作。各单位应按计划落实任务。",
     "（一）统筹推进重点项目建设和公共服务保障工作。"),
])
def test_h2_inline_heading_boundaries(text, heading):
    item = classify(["标题", text])[1]
    assert item.role == Role.H2
    assert item.text[:item.lead_len] == heading
    assert item.confidence == "low"  # 纯文字无法确定短句一定是标题，保留确认入口
    assert not item.confirmed
    assert item.lead_len == lead_len(Role.H2, item.text)


@pytest.mark.parametrize("text", [
    "（一）工作目标", "（一）工作目标：", "（一）10月1日-10月7日，筹备阶段",
    "（一）统筹推进重点项目建设和公共服务保障工作",
])
def test_h2_standalone_headings_are_preserved(text):
    item = classify(["标题", text])[1]
    assert item.role == Role.H2
    assert item.lead_len == len(text)


def test_numbered_body_is_not_reclassified_by_position():
    items = classify(["（一）", "（二）", "（三）", "正文。", "（四）", "2026年10月1日"])
    assert [it.role for it in items] == [Role.BODY] * 5 + [Role.DATE]


def test_numbered_body_still_normalizes_prefix_and_allows_manual_override():
    item = classify(["标题", "(一)、请按时提交材料。"])[1]
    assert item.role == Role.BODY
    assert item.text == "（一）请按时提交材料。"
    assert item.fixed == "“(一)、”→“（一）”"
    # 用户明确选择二级标题时，整段仍可使用标题格式；切回正文时不残留标题范围。
    assert lead_len(Role.H2, item.text) == len(item.text)
    assert lead_len(Role.BODY, item.text) == 0


def _font_and_bold(run):
    return (run._r.rPr.rFonts.get(qn("w:eastAsia")), bool(run.bold))


@pytest.mark.parametrize("requirements, body_font, heading_font", [
    ("", "仿宋_GB2312", "楷体_GB2312"),
    ("正文（宋体，小四，不加粗，顶格）\n二级标题（黑体，三号，加粗，首行缩进2字符）", "宋体", "黑体"),
])
def test_numbered_body_word_formatting(tmp_path, requirements, body_font, heading_font):
    src, out = tmp_path / "原稿.docx", tmp_path / "排版.docx"
    lines = ["部门三年发展规划", "一、工作要求", "（一）工作目标", "正文内容。",
             "（二）工作安排：各科室按时提交材料。", "二、结语", *CONCLUSION]
    doc = Document()
    for text in lines:
        run = doc.add_paragraph(text).runs[0]
        run.font.name, run.bold = "微软雅黑", True  # 原稿格式不参与判断
    doc.save(src)
    original = src.read_bytes()
    items, warnings = format_file(str(src), str(out), requirements)
    assert src.read_bytes() == original
    assert warnings == []
    paragraphs = Document(out).paragraphs
    assert [p.text for p in paragraphs] == lines
    for paragraph in paragraphs[-2:]:
        assert all(_font_and_bold(r) == (body_font, False) for r in paragraph.runs)
        indent = paragraph._p.pPr.find(qn("w:ind"))
        assert indent.get(qn("w:firstLineChars")) == ("200" if not requirements else "0")
    assert _font_and_bold(paragraphs[2].runs[0]) == (heading_font, True)
    inline = paragraphs[4].runs
    assert inline[0].text == "（二）工作安排："
    assert _font_and_bold(inline[0]) == (heading_font, True)
    assert _font_and_bold(inline[1]) == (body_font, False)
    assert all(item.role == Role.BODY for item in items[-2:])
